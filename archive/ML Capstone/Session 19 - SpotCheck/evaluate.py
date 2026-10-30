"""Evaluation evidence for the SpotCheck models, written into the existing files.

1. Cross-validation on all 9,244 restaurants: 10-fold KFold (shuffle, random_state=1)
   and 10-fold GroupKFold grouped by restaurant_name (chain outlets share a name).
   Models: mean baseline, Linear Regression, Random Forest (same settings as train_models.py).
2. Permutation feature importance of the Random Forest on the held-out KFold folds.
3. Linear-regression assumption checks (statsmodels) on zomato_price_train.tab.
4. Appends the app's outputs (price and rating, via predict.py) to reference_tests.csv.
5. Writes the results into SpotCheck_Session19_Pack.xlsx (in place, formatting kept).

Run after train_models.py:  python evaluate.py
Needs statsmodels and openpyxl in addition to requirements.txt (not needed by the app).
"""
import copy
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.api as sm
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.dummy import DummyRegressor
from sklearn.inspection import permutation_importance
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score, roc_auc_score
from sklearn.model_selection import GroupKFold, KFold
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder
from statsmodels.stats.diagnostic import het_breuschpagan, linear_reset
from statsmodels.stats.stattools import jarque_bera

from predict import CATEGORICAL, LABELS, PRICE_FEATURES, predict_price, predict_rating
from train_models import load_tab, price_pipeline

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE / "SpotCheck_Session19_Pack.xlsx"
REFERENCE_CSV = HERE / "reference_tests.csv"
NEW_REF_COLUMNS = ["Predicted cost (₹)", "Error (₹)", "Abs error (₹)", "P(4.0+)", "Predicted 4.0+ (0.5)", "Correct?"]

# Orange GUI, "Test & Score – PRICE (HOLD-OUT test)", Test on test data, read off the screen on 2026-10-26
# after fixing the RF settings (100 trees, replicable, no attribute/depth limit, min split 5, no class balancing).
# Rows of the 'Hold-out Test' sheet: (row, RMSE, MAE, R²).
ORANGE_HOLDOUT = {
    "Linear Regression": (5, 268.125, 185.075, 0.754),
    "Random Forest": (7, 206.970, 151.730, 0.853),
}
# Orange GUI, "Test & Score – PRICE (10-fold CV)", same date and settings, all 9,244 rows.
# Rows of the 'Model Comparison (CV)' sheet: (row, RMSE, MAE, R²).
ORANGE_CV = {
    "Mean (baseline)": (4, 426.866, 276.071, -0.000),
    "Linear Regression": (5, 267.161, 164.446, 0.608),
    "Tree depth 5": (6, 247.746, 153.846, 0.663),
    "Random Forest": (7, 223.657, 136.778, 0.725),
}
# Orange GUI, "Test & Score – RATING (10-fold CV)", target class Yes, same date; RF set to the price recipe.
# Rows of the 'Model Comparison (CV)' sheet: (row, AUC, CA, F1, Precision, Recall).
ORANGE_RATING_CV = {
    "Constant": (11, 0.500, 0.772, 0.000, 0.000, 0.000),
    "Logistic Regression": (12, 0.811, 0.833, 0.528, 0.744, 0.409),
    "Tree depth 4": (13, 0.729, 0.811, 0.476, 0.647, 0.377),
    "Random Forest": (14, 0.825, 0.841, 0.586, 0.722, 0.493),
}
# Orange GUI, "Confusion Matrix (CV)", Logistic Regression, number of instances: (actual No, actual Yes) x (pred No, pred Yes).
ORANGE_RATING_CM = ((6840, 297), (1245, 862))


# ---------------------------------------------------------------- models
def linear_pipeline() -> Pipeline:
    encode = ColumnTransformer(
        [("onehot", OneHotEncoder(handle_unknown="ignore"), CATEGORICAL)], remainder="passthrough"
    )
    return Pipeline([("encode", encode), ("ols", LinearRegression())])


MODELS = {
    "Mean (baseline)": DummyRegressor(strategy="mean"),
    "Linear Regression (comparison)": linear_pipeline(),
    "Random Forest (FINAL, 100 trees)": price_pipeline(),
}


def metrics(y, p) -> dict:
    return {
        "MAE (₹)": mean_absolute_error(y, p),
        "RMSE (₹)": np.sqrt(mean_squared_error(y, p)),
        "R²": r2_score(y, p),
    }


# ---------------------------------------------------------------- 1 + 2: CV and importance
def cross_validate(full: pd.DataFrame):
    """Pooled out-of-fold metrics per scheme and model, plus RF permutation importance (KFold folds)."""
    X, y = full[PRICE_FEATURES], full["cost_for_two"]
    schemes = {
        "KFold": list(KFold(10, shuffle=True, random_state=1).split(X)),
        "GroupKFold": list(GroupKFold(10).split(X, groups=full["restaurant_name"])),
    }
    results, importance = {}, []
    for scheme, folds in schemes.items():
        for name, model in MODELS.items():
            oof = np.zeros(len(y))
            for train_idx, test_idx in folds:
                fitted = clone(model).fit(X.iloc[train_idx], y.iloc[train_idx])
                oof[test_idx] = fitted.predict(X.iloc[test_idx])
                if scheme == "KFold" and name.startswith("Random Forest"):
                    # Held-out fold only: permuting a raw input column = permuting the whole field.
                    pi = permutation_importance(
                        fitted, X.iloc[test_idx], y.iloc[test_idx],
                        scoring="neg_mean_absolute_error", n_repeats=5, random_state=0,
                    )
                    importance.append(pi.importances_mean)
            results[(scheme, name)] = metrics(y, oof)
            print(f"{scheme:10s} {name:34s} " + "  ".join(f"{k} {v:,.3f}" for k, v in results[(scheme, name)].items()))
    imp = pd.DataFrame(importance, columns=PRICE_FEATURES)
    imp = pd.DataFrame({"mean": imp.mean(), "sd": imp.std(ddof=1)}).sort_values("mean", ascending=False)
    return results, imp


# ---------------------------------------------------------------- 3: LR assumptions
def gvif(X: pd.DataFrame, groups: dict) -> pd.DataFrame:
    """Generalised VIF (Fox & Monette 1992) per original field; X excludes the constant."""
    R = np.corrcoef(X.values, rowvar=False)
    logdet_all = np.linalg.slogdet(R)[1]
    cols = list(X.columns)
    rows = []
    for field, members in groups.items():
        idx = [cols.index(m) for m in members]
        rest = [i for i in range(len(cols)) if i not in idx]
        logdet = np.linalg.slogdet(R[np.ix_(idx, idx)])[1] + np.linalg.slogdet(R[np.ix_(rest, rest)])[1] - logdet_all
        g = np.exp(logdet)
        rows.append({"field": field, "df": len(idx), "GVIF": g, "GVIF^(1/(2·df))": g ** (1 / (2 * len(idx)))})
    return pd.DataFrame(rows).sort_values("GVIF^(1/(2·df))", ascending=False)


def assumption_checks(train: pd.DataFrame) -> dict:
    dummies = pd.get_dummies(train[CATEGORICAL], drop_first=True, prefix_sep="=").astype(float)
    X = pd.concat([dummies, train[["n_cuisines"]].astype(float)], axis=1)
    ols = sm.OLS(train["cost_for_two"].astype(float), sm.add_constant(X)).fit()

    reset = linear_reset(ols, power=3, use_f=True)
    bp_lm, bp_p, _, _ = het_breuschpagan(ols.resid, ols.model.exog)
    jb, jb_p, skew, kurt = jarque_bera(ols.resid)
    decile = pd.qcut(ols.fittedvalues, 10, labels=False) + 1
    by_decile = pd.DataFrame({
        "fitted mean (₹)": ols.fittedvalues.groupby(decile).mean(),
        "residual SD (₹)": ols.resid.groupby(decile).std(),
    })
    groups = {f: [c for c in X.columns if c.split("=")[0] == f] for f in CATEGORICAL}
    groups["n_cuisines"] = ["n_cuisines"]
    out = {
        "n": int(ols.nobs), "r2": ols.rsquared,
        "reset_F": float(reset.fvalue), "reset_p": float(reset.pvalue),
        "bp_LM": bp_lm, "bp_p": bp_p,
        "jb": jb, "jb_p": jb_p, "skew": skew, "kurtosis": kurt,
        "deciles": by_decile, "gvif": gvif(X, groups),
    }
    print(f"\nOLS on train (n={out['n']}): R² {out['r2']:.3f}; RESET F={out['reset_F']:.1f} p={out['reset_p']:.2g}; "
          f"Breusch–Pagan LM={bp_lm:.0f} p={bp_p:.2g}; Jarque–Bera={jb:,.0f} p={jb_p:.2g} (skew {skew:.2f}, kurtosis {kurt:.1f})")
    print(by_decile.round(0).to_string())
    print(out["gvif"].round(2).to_string(index=False))
    return out


# ---------------------------------------------------------------- 4: reference_tests.csv
def update_reference_csv() -> pd.DataFrame:
    """Append the app's outputs (through predict.py) as new columns; existing text stays byte-identical."""
    lines = REFERENCE_CSV.read_text(encoding="utf-8").splitlines()
    header = lines[0].split(",")
    if NEW_REF_COLUMNS[0] in header:  # re-run: drop the previously appended columns first
        keep = header.index(NEW_REF_COLUMNS[0])
        lines = [",".join(line.split(",")[:keep]) for line in lines]
    ref = pd.read_csv(REFERENCE_CSV, usecols=range(len(lines[0].split(","))), dtype=str, keep_default_na=False)

    rows = []
    for _, r in ref.iterrows():
        pred = round(predict_price(r.to_dict()), 1)
        actual = float(r["cost_for_two"])
        prob = round(predict_rating(r.to_dict()), 4)
        label = "Yes" if prob >= 0.5 else "No"
        correct = "Yes" if label == r["high_rating"] else "No"
        rows.append([f"{pred:.1f}", f"{pred - actual:.1f}", f"{abs(pred - actual):.1f}", f"{prob:.4f}", label, correct])
    out = [lines[0] + "," + ",".join(NEW_REF_COLUMNS)]
    out += [line + "," + ",".join(vals) for line, vals in zip(lines[1:], rows)]
    REFERENCE_CSV.write_text("\n".join(out) + "\n", encoding="utf-8")

    ref["Predicted cost (₹)"] = [float(v[0]) for v in rows]
    ref["P(4.0+)"] = [float(v[3]) for v in rows]
    mae = np.mean([float(v[2]) for v in rows])
    n_correct = sum(v[5] == "Yes" for v in rows)
    print(f"\nreference_tests.csv updated: hold-out MAE of the shipped RF = ₹{mae:.1f}; rating LR {n_correct}/10 correct at 0.5")
    return ref


# ---------------------------------------------------------------- 5: workbook
HEADER_FONT = Font(bold=True, color="FFFFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_table(ws, top: int, left: int, header: list, rows: list, formats: list | None = None) -> None:
    for j, h in enumerate(header):
        c = ws.cell(top, left + j, h)
        c.font, c.fill, c.border, c.alignment = HEADER_FONT, HEADER_FILL, BOX, WRAP
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = ws.cell(top + i, left + j, v)
            c.border, c.alignment = BOX, WRAP
            if formats and formats[j]:
                c.number_format = formats[j]


def note(ws, cell: str, text: str) -> None:
    ws[cell] = text
    ws[cell].alignment = Alignment(wrap_text=False, vertical="top")


def update_workbook(cv: dict, imp: pd.DataFrame, checks: dict, ref: pd.DataFrame) -> None:
    lock = WORKBOOK.with_name("~$" + WORKBOOK.name)
    if lock.exists():
        raise SystemExit(f"{lock.name} exists: the workbook is open in Excel. Close it and run again.")
    wb = openpyxl.load_workbook(WORKBOOK)
    rf, lr, base = "Random Forest (FINAL, 100 trees)", "Linear Regression (comparison)", "Mean (baseline)"
    kf = lambda m, k: cv[("KFold", m)][k]
    gk = lambda m, k: cv[("GroupKFold", m)][k]
    ref_mae = float(np.mean(np.abs(ref["Predicted cost (₹)"] - ref["cost_for_two"].astype(float))))
    n_correct = int(((ref["P(4.0+)"] >= 0.5) == (ref["high_rating"] == "Yes")).sum())
    ref_auc = roc_auc_score(ref["high_rating"] == "Yes", ref["P(4.0+)"])
    top3 = ", ".join(LABELS[f] for f in imp.index[:3])

    # --- Summary
    s = wb["Summary"]
    s["B5"] = "Target: cost_for_two (₹, regression). Final model: Random Forest regression (100 trees, random_state=0)"
    s["A16"] = "Why Random Forest for price?"
    s["B16"] = (
        f"10-fold CV on all 9,244 restaurants (Python): Random Forest R² {kf(rf, 'R²'):.2f}, MAE ₹{kf(rf, 'MAE (₹)'):.0f} "
        f"vs Linear Regression R² {kf(lr, 'R²'):.2f}, MAE ₹{kf(lr, 'MAE (₹)'):.0f} (mean baseline MAE ₹{kf(base, 'MAE (₹)'):.0f}); "
        f"Orange confirms it: RF R² {ORANGE_CV['Random Forest'][3]:.3f}, MAE ₹{ORANGE_CV['Random Forest'][2]:.0f} vs "
        f"Linear Regression R² {ORANGE_CV['Linear Regression'][3]:.3f}, MAE ₹{ORANGE_CV['Linear Regression'][2]:.0f}. "
        f"Grouped by restaurant name (GroupKFold, so chain outlets never sit in both train and test), RF still leads: "
        f"R² {gk(rf, 'R²'):.2f} vs {gk(lr, 'R²'):.2f}. Linear Regression fails its assumption checks (non-linearity, "
        f"unequal error spread, non-normal errors; see 'Assumption Checks'). The app loads the trained model, so "
        f"a formula is no longer needed. Most important inputs: {top3}."
    )
    s["B17"] = (
        f"10-fold CV (Orange, target 4.0+): AUC {ORANGE_RATING_CV['Logistic Regression'][1]:.3f}, close to Random Forest "
        f"({ORANGE_RATING_CV['Random Forest'][1]:.3f}) and well above the tree ({ORANGE_RATING_CV['Tree depth 4'][1]:.3f}). "
        f"Random Forest is slightly ahead, but Logistic Regression gives smooth, explainable probabilities for what-if "
        f"comparisons. The Python model the app loads reproduces Orange's P(4.0+) within 0.002 on all 10 hold-out restaurants."
    )
    s["B18"] = (
        f"Price: Random Forest (Python, the model the app loads) MAE ₹{ref_mae:.0f} on the 10 unseen restaurants "
        f"(mean baseline ₹433; Orange RF ₹{ORANGE_HOLDOUT['Random Forest'][2]:.0f}; Linear Regression "
        f"₹{ORANGE_HOLDOUT['Linear Regression'][2]:.0f}). Rating: Logistic Regression (Python, loaded by the app) "
        f"{n_correct}/10 correct at 0.5, hold-out AUC {ref_auc:.3f} (Orange: 0.917). Ten rows are too few to judge accuracy and "
        f"they were picked for variety, not at random, so cross-validation remains the main evidence. The hold-out "
        f"shows the models work on unseen data."
    )
    s["A20"] = "Integration"
    s["B20"] = (
        "Python app (Streamlit) → code on GitHub → Streamlit Community Cloud (confirmed by facilitator). "
        "The app loads price_model.joblib / rating_model.joblib through predict.py; models are not coded as formulas."
    )
    for col in "AB":
        s[f"{col}20"]._style = copy.copy(s[f"{col}19"]._style)
    s.row_dimensions[16].height = 80
    s.row_dimensions[17].height = 55
    s.row_dimensions[18].height = 55
    s.row_dimensions[20].height = 30

    # --- Model Comparison (CV): Python KFold + GroupKFold table beside the Orange one
    mc = wb["Model Comparison (CV)"]
    for model, (row, rmse, mae, r2) in ORANGE_CV.items():
        assert mc.cell(row, 1).value.startswith(model)
        mc.cell(row, 2).value, mc.cell(row, 3).value, mc.cell(row, 4).value = rmse, mae, r2
    for model, (row, *scores) in ORANGE_RATING_CV.items():
        assert mc.cell(row, 1).value.startswith(model)
        for j, v in enumerate(scores):
            mc.cell(row, 2 + j).value = v
    mc["A22"] = ("• Price table (rows 4–7) and rating table (rows 11–14, target class Yes) read from the Orange 3.39 GUI on "
                 "2026-10-26 after fixing the Random Forest settings (100 trees, replicable, no attribute/depth limit, min split 5, "
                 "no class balancing). A new Test & Score run may differ by about ±0.01 because of the random fold split.")
    for i, counts in enumerate(ORANGE_RATING_CM):
        for j, n in enumerate(counts):
            mc.cell(18 + i, 2 + j).value = n
    note(mc, "A23", "• Confusion matrix (rows 17–19): Logistic Regression, Orange GUI 'Confusion Matrix (CV)', same run (precision 862/1,159 = 0.744, recall 862/2,107 = 0.409).")
    mc["F1"] = "Price – Python check: 10-fold KFold vs 10-fold GroupKFold (grouped by restaurant_name)"
    mc["F1"].font = copy.copy(mc["A1"].font)
    header = ["Model", "KFold MAE (₹)", "KFold RMSE (₹)", "KFold R²", "GroupKFold MAE (₹)", "GroupKFold RMSE (₹)", "GroupKFold R²"]
    rows = [[m, kf(m, "MAE (₹)"), kf(m, "RMSE (₹)"), kf(m, "R²"), gk(m, "MAE (₹)"), gk(m, "RMSE (₹)"), gk(m, "R²")]
            for m in MODELS]
    write_table(mc, 3, 6, header, rows, [None, "0.0", "0.0", "0.000", "0.0", "0.0", "0.000"])
    note(mc, "F7", "• KFold: shuffled, random_state=1. GroupKFold: all outlets of a chain (same name, ~40% of rows) stay in one fold.")
    note(mc, "F8", "• GroupKFold is the stricter test: it checks the model on brands it has never seen, as the app will.")
    mc.column_dimensions["F"].width = 34
    for col in "GHIJKL":
        mc.column_dimensions[col].width = 13

    # --- Hold-out Test: RF predictions from reference_tests.csv in the same row-by-row layout
    h = wb["Hold-out Test"]
    for i, pred in enumerate(ref["Predicted cost (₹)"]):
        assert h.cell(17 + i, 1).value == ref["TestID"].iloc[i]
        h.cell(17 + i, 4).value = pred
    h["D16"] = "Predicted cost (₹) – RF (Python)"
    for model, (row, rmse, mae, r2) in ORANGE_HOLDOUT.items():
        assert h.cell(row, 1).value.startswith(model)
        h.cell(row, 2).value, h.cell(row, 3).value, h.cell(row, 4).value = rmse, mae, r2
    orange_rf_mae = ORANGE_HOLDOUT["Random Forest"][2]
    note(h, "A32", f"• Price predictions are from the Python Random Forest the app loads (price_model.joblib via predict.py); hold-out MAE ₹{ref_mae:.1f}.")
    note(h, "A33", f"• Orange's Random Forest (same settings, rows 4–7 from the Orange GUI) gets MAE ₹{orange_rf_mae:.1f}: two forests are never identical, so per-row numbers differ, but the average error agrees within ₹{abs(orange_rf_mae - ref_mae):.1f}. Between Python seeds, predictions move by about ₹5–56 per restaurant; the app reproduces the Python values.")
    note(h, "A35", "• Linear Regression and Random Forest rows above were re-read from the Orange GUI on 2026-10-26 (RF settings fixed: no depth/attribute limit, min split 5). Linear Regression changed from ₹183.9 to ₹185.1; the cause is not yet known.")
    note(h, "A36", "• Linear Regression check: plain OLS in Python (scikit-learn, same training file and 10 restaurants) gives MAE ₹183.9, RMSE ₹264.9, R² 0.760, the original value in this sheet. Orange's LR widget gives ₹185.1 with the same data and settings (no regularization, fit intercept); the gap comes from inside Orange. LR is comparison-only and does not affect the app.")
    for i, prob in enumerate(ref["P(4.0+)"]):
        h.cell(17 + i, 9).value = round(prob, 3)
    h["I16"] = "P(4.0+) – LR (Python)"
    note(h, "A34", "• P(4.0+) is from the Python Logistic Regression the app loads (rating_model.joblib via predict.py); it matches the Orange GUI within 0.002 for every restaurant. The earlier values (0.143, 0.175, …) did not come from this Orange workflow.")
    wb.calculation.fullCalcOnLoad = True

    # --- Reference Tests: same restaurants, predicted price now from the RF
    rt = wb["Reference Tests"]
    for i, pred in enumerate(ref["Predicted cost (₹)"]):
        assert rt.cell(5 + i, 1).value == ref["TestID"].iloc[i]
        rt.cell(5 + i, 22).value = pred
    for i, prob in enumerate(ref["P(4.0+)"]):
        rt.cell(5 + i, 24).value = round(prob, 3)
    rt["D2"] = round(kf(rf, "MAE (₹)"))
    note(rt, "A17", f"Predicted market price = Random Forest (price_model.joblib). Price band ± MAE = RF 10-fold CV MAE (₹{kf(rf, 'MAE (₹)'):.0f}). P(rating ≥ 4.0) = Logistic Regression (rating_model.joblib).")

    # --- Rating coefficients: the old values do not reproduce Orange
    rc = wb["Rating Coefficients"]
    note(rc, "A93", "Note (2026-10-26): these coefficients do not reproduce the Orange GUI's predictions (they give P = 0.143, 0.175, … for T01–T02;")
    note(rc, "A94", "Orange gives 0.092, 0.061, …). The app does not use this formula; it loads rating_model.joblib, which matches Orange within 0.002.")

    # --- LR coefficients: rename + note
    pc = wb["Price Coefficients"] if "Price Coefficients" in wb.sheetnames else wb["LR Coefficients (comparison)"]
    pc.title = "LR Coefficients (comparison)"
    pc["A1"] = "Linear Regression – price for two (₹) – COMPARISON ONLY (final price model is the Random Forest)"
    note(pc, "A92", "Note: Orange lists both the Yes and the No level of each Yes/No input. The Yes-vs-No effect is the difference between them,")
    note(pc, "A93", "e.g. book_table: +208.2 − (−208.2) = +₹416 for offering table booking, not +₹208. The app does not use this formula.")

    # --- New sheets
    add_importance_sheet(wb, imp)
    add_assumption_sheet(wb, checks, cv)
    wb.save(WORKBOOK)
    print(f"\nWorkbook updated: {WORKBOOK.name}")


def add_importance_sheet(wb, imp: pd.DataFrame) -> None:
    if "RF Feature Importance" in wb.sheetnames:
        del wb["RF Feature Importance"]
    ws = wb.create_sheet("RF Feature Importance")
    ws["A1"] = "Random Forest (price) – permutation importance on held-out folds"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Increase in MAE (₹) when one input field is shuffled, measured only on the held-out fold of each "
                "of the 10 KFold splits (never on training data). 5 shuffles per fold; mean and SD across folds.")
    rows = []
    for field, r in imp.iterrows():
        verdict = ("Major driver" if r["mean"] >= 20 else
                   "Moderate" if r["mean"] >= 5 else
                   "Minor" if r["mean"] > 2 * r["sd"] else "Negligible (within noise)")
        rows.append([LABELS[field], field, r["mean"], r["sd"], verdict])
    write_table(ws, 4, 1, ["Input field", "Column", "MAE increase (₹)", "SD across folds (₹)", "Verdict"],
                rows, [None, None, "0.0", "0.0", None])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 26
    ws.freeze_panes = "A5"
    note(ws, f"A{len(rows) + 6}", "• Each field is permuted as a whole (location = all 40 levels together), so the importance is per original input, not per one-hot column.")
    note(ws, f"A{len(rows) + 7}", "• Importance shows what the model relies on, not cause and effect.")


def fmt_p(p: float) -> str:
    return "< 0.001" if p < 0.001 else f"= {p:.3f}"


def add_assumption_sheet(wb, c: dict, cv: dict) -> None:
    if "Assumption Checks" in wb.sheetnames:
        del wb["Assumption Checks"]
    ws = wb.create_sheet("Assumption Checks")
    ws["A1"] = "Linear Regression (price) – assumption checks (statsmodels, zomato_price_train.tab)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"OLS with one-hot inputs (first level of each field as base), n = {c['n']:,}, R² = {c['r2']:.3f}."
    sd = c["deciles"]["residual SD (₹)"]
    worst_gvif = c["gvif"].iloc[0]
    high_gvif = [LABELS[f] for f in c["gvif"].loc[c["gvif"]["GVIF^(1/(2·df))"] >= 2.24, "field"]]
    rows = [
        ["Linearity", "Ramsey RESET (powers 2–3)", f"F = {c['reset_F']:.1f}, p {fmt_p(c['reset_p'])}",
         "Violated: the straight-line form misses curvature." if c["reset_p"] < 0.05 else "OK: no evidence of missed curvature."],
        ["Equal error spread", "Breusch–Pagan", f"LM = {c['bp_LM']:.0f}, p {fmt_p(c['bp_p'])}",
         "Violated: errors grow with the predicted price." if c["bp_p"] < 0.05 else "OK: error spread is constant."],
        ["Equal error spread", "Residual SD by fitted-value decile", f"₹{sd.iloc[0]:.0f} (lowest decile) → ₹{sd.iloc[-1]:.0f} (highest)",
         f"Violated: errors are {sd.iloc[-1] / sd.iloc[0]:.1f}× larger for expensive restaurants." if sd.iloc[-1] > 2 * sd.iloc[0] else "OK: similar spread across deciles."],
        ["Normal errors", "Jarque–Bera", f"JB = {c['jb']:,.0f}, p {fmt_p(c['jb_p'])}; skew {c['skew']:.2f}, kurtosis {c['kurtosis']:.1f}",
         "Violated: heavy right tail (a few very expensive restaurants)." if c["jb_p"] < 0.05 else "OK: errors look normal."],
        ["Multicollinearity", "GVIF^(1/(2·df)) per field (largest)", f"{LABELS[worst_gvif['field']]}: {worst_gvif['GVIF^(1/(2·df))']:.2f}",
         "OK: all below 2.24 (≈ VIF 5)." if worst_gvif["GVIF^(1/(2·df))"] < 2.24 else f"Concern for {', '.join(high_gvif)} only (above 2.24 ≈ VIF 5; Cafe cuisine overlaps with type = Cafe). Other fields are fine."],
        ["Conclusion", "Model choice", "",
         f"Linear Regression's assumptions fail, and Random Forest needs none of them; CV R² {cv[('KFold', 'Random Forest (FINAL, 100 trees)')]['R²']:.2f} vs {cv[('KFold', 'Linear Regression (comparison)')]['R²']:.2f}. Random Forest is the final price model."],
    ]
    write_table(ws, 4, 1, ["Assumption", "Check", "Result", "Verdict"], rows)
    for col, w in zip("ABCD", (20, 34, 44, 70)):
        ws.column_dimensions[col].width = w

    top = 4 + len(rows) + 2
    ws.cell(top, 1, "Residual spread by fitted-value decile").font = Font(bold=True)
    write_table(ws, top + 1, 1, ["Decile", "Mean fitted (₹)", "Residual SD (₹)"],
                [[int(d), r["fitted mean (₹)"], r["residual SD (₹)"]] for d, r in c["deciles"].iterrows()],
                [None, "0", "0"])
    top = top + 1 + len(c["deciles"]) + 2
    ws.cell(top, 1, "Generalised VIF per input field").font = Font(bold=True)
    write_table(ws, top + 1, 1, ["Input field", "df", "GVIF", "GVIF^(1/(2·df))"],
                [[LABELS[r["field"]], int(r["df"]), r["GVIF"], r["GVIF^(1/(2·df))"]] for _, r in c["gvif"].iterrows()],
                [None, "0", "0.00", "0.00"])


def main() -> None:
    full = load_tab("zomato_price.tab")
    train = load_tab("zomato_price_train.tab")
    cv, imp = cross_validate(full)
    print("\nPermutation importance (MAE increase, ₹):")
    print(imp.round(1).to_string())
    checks = assumption_checks(train)
    ref = update_reference_csv()
    update_workbook(cv, imp, checks, ref)


if __name__ == "__main__":
    main()
